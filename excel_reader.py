from openpyxl import load_workbook


def read_participants(file_path):
    workbook = load_workbook(file_path)
    worksheet = workbook.active

    participants = []

    # Get column headers from the first row
    headers = {}

    for cell in worksheet[1]:
        if cell.value:
            headers[str(cell.value).strip().lower()] = cell.column

    # Make sure required columns exist
    if "name" not in headers:
        raise ValueError("Excel file does not contain a 'Name' column.")

    if "email" not in headers:
        raise ValueError("Excel file does not contain an 'Email' column.")

    name_column = headers["name"]
    email_column = headers["email"]

    # Read participant data starting from row 2
    for row in range(2, worksheet.max_row + 1):
        name = worksheet.cell(row=row, column=name_column).value
        email = worksheet.cell(row=row, column=email_column).value

        if name is None and email is None:
            continue

        participant = {
            "name": str(name).strip() if name else "",
            "email": str(email).strip() if email else "",
        }

        participants.append(participant)

    return participants